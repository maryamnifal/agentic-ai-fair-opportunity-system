import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
        # Header
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("FairWork AI — Technical Report | IT 3041")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 120, 120)
        
        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Faculty of Computing — Sri Lanka Institute of Information Technology (SLIIT)")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(140, 140, 140)

def apply_custom_styles(doc):
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(40, 40, 40)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

def format_inline_runs(paragraph, text, base_font_size=11, is_italic=False):
    # Regex parser for bold, italic, inline code
    # Matches: `code`, **bold**, *italic*
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) >= 4:
            run = paragraph.add_run(tok[2:-2])
            run.font.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            if is_italic:
                run.font.italic = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) >= 2:
            run = paragraph.add_run(tok[1:-1])
            run.font.italic = True
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) >= 2:
            run = paragraph.add_run(tok[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 1)
            run.font.color.rgb = RGBColor(180, 40, 40)
        else:
            run = paragraph.add_run(tok)
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            if is_italic:
                run.font.italic = True

def convert_markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    add_header_footer(doc)
    apply_custom_styles(doc)

    in_code_block = False
    code_lines = []
    code_language = ""

    in_table = False
    table_lines = []

    def flush_table(lines_batch):
        if not lines_batch:
            return
        # Parse table rows
        rows_data = []
        for l in lines_batch:
            # Check separator row
            if re.match(r'^\s*\|?\s*[-:]+\s*\|', l):
                continue
            cols = [c.strip() for c in l.strip().strip('|').split('|')]
            if cols:
                rows_data.append(cols)
        
        if not rows_data:
            return

        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for r_idx, row in enumerate(rows_data):
            for c_idx in range(num_cols):
                cell = table.cell(r_idx, c_idx)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                
                cell_text = row[c_idx] if c_idx < len(row) else ""
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(2)
                cp.paragraph_format.space_before = Pt(2)
                
                if r_idx == 0:
                    set_cell_background(cell, "1F4E79") # Deep Blue
                    format_inline_runs(cp, cell_text, base_font_size=9.5)
                    for r in cp.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F2F5F9") # Soft Ice Blue
                    else:
                        set_cell_background(cell, "FFFFFF")
                    format_inline_runs(cp, cell_text, base_font_size=9)
        
        # Space after table
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(6)

    def flush_code(lines_batch, lang):
        if not lines_batch:
            return
        full_code = "".join(lines_batch).rstrip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        
        # Border and shading
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'''
            <w:pBdr {nsdecls("w")}>
                <w:left w:val="single" w:sz="18" w:space="8" w:color="005A9E"/>
            </w:pBdr>
        ''')
        pPr.append(pBdr)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F6F9"/>')
        pPr.append(shd)

        run = p.add_run(full_code)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(45, 55, 72)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code(code_lines, code_language)
                code_lines = []
                code_language = ""
            else:
                if in_table:
                    in_table = False
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
                code_language = stripped[3:].strip()
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle tables
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                in_table = True
                table_lines = [line]
            else:
                table_lines.append(line)
            i += 1
            continue
        elif in_table:
            in_table = False
            flush_table(table_lines)
            table_lines = []

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Horizontal rules
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'''
                <w:pBdr {nsdecls("w")}>
                    <w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/>
                </w:pBdr>
            ''')
            pPr.append(pBdr)
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(8)
            h.paragraph_format.keep_with_next = True
            format_inline_runs(h, stripped[2:].strip(), base_font_size=20)
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.bold = True
                r.font.color.rgb = RGBColor(16, 44, 87)
            i += 1
            continue
        elif stripped.startswith("## "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            format_inline_runs(h, stripped[3:].strip(), base_font_size=15)
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.bold = True
                r.font.color.rgb = RGBColor(31, 78, 121)
            i += 1
            continue
        elif stripped.startswith("### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            format_inline_runs(h, stripped[4:].strip(), base_font_size=12.5)
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.bold = True
                r.font.color.rgb = RGBColor(46, 117, 182)
            i += 1
            continue
        elif stripped.startswith("#### "):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            format_inline_runs(h, stripped[5:].strip(), base_font_size=11.5)
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.bold = True
                r.font.color.rgb = RGBColor(60, 60, 60)
            i += 1
            continue

        # Blockquotes
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'''
                <w:pBdr {nsdecls("w")}>
                    <w:left w:val="single" w:sz="18" w:space="8" w:color="2E75B6"/>
                </w:pBdr>
            ''')
            pPr.append(pBdr)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F7FA"/>')
            pPr.append(shd)
            text_inside = stripped.lstrip('>').strip()
            format_inline_runs(p, text_inside, base_font_size=10.5, is_italic=True)
            i += 1
            continue

        # Bullet lists
        if re.match(r'^[-*+]\s+', stripped):
            item_text = re.sub(r'^[-*+]\s+', '', stripped)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            format_inline_runs(p, item_text, base_font_size=11)
            i += 1
            continue

        # Numbered lists
        if re.match(r'^\d+\.\s+', stripped):
            item_text = re.sub(r'^\d+\.\s+', '', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            format_inline_runs(p, item_text, base_font_size=11)
            i += 1
            continue

        # Figure placeholder callout
        if stripped.startswith("[Insert Screenshot Here:") or stripped.startswith("[Insert screenshot here"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'''
                <w:pBdr {nsdecls("w")}>
                    <w:top w:val="dashed" w:sz="12" w:space="8" w:color="0078D4"/>
                    <w:bottom w:val="dashed" w:sz="12" w:space="8" w:color="0078D4"/>
                    <w:left w:val="dashed" w:sz="12" w:space="8" w:color="0078D4"/>
                    <w:right w:val="dashed" w:sz="12" w:space="8" w:color="0078D4"/>
                </w:pBdr>
            ''')
            pPr.append(pBdr)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F6FF"/>')
            pPr.append(shd)

            run = p.add_run(f"\n📷 {stripped}\n(Paste application screenshot into this placeholder box)\n")
            run.font.name = 'Calibri'
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 90, 180)
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        format_inline_runs(p, stripped, base_font_size=11)
        i += 1

    # Cleanup trailing table
    if in_table:
        flush_table(table_lines)

    doc.save(docx_path)
    print(f"Successfully generated Word document at: {docx_path}")

if __name__ == '__main__':
    md_in = r'c:\Users\marya\agentic-ai-fair-opportunity-system\docs\FairWork_AI_Technical_Report.md'
    docx_out = r'c:\Users\marya\agentic-ai-fair-opportunity-system\docs\FairWork_AI_Technical_Report.docx'
    convert_markdown_to_docx(md_in, docx_out)
